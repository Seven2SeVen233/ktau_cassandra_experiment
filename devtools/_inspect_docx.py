"""检查 docx 表格内容与全文引用编号。"""
import sys
import docx

path = r"C:\Users\34333\OneDrive\Desktop\Kendall-tau\Kendall_Tau_Log_Reconciliation_ICSDC_prepare.docx"
d = docx.Document(path)

print("=== 引用编号出现位置 ===")
import re
for i, p in enumerate(d.paragraphs):
    t = p.text
    for m in re.finditer(r"\[(\d{1,2})\]", t):
        print(f"para[{i}] [{m.group(1)}]  ...{t[max(0,m.start()-40):m.end()+40]}...")

print("\n=== 全部表格内容 ===")
for ti, tb in enumerate(d.tables):
    print(f"\n--- Table {ti}: {len(tb.rows)}x{len(tb.columns)} ---")
    for ri, row in enumerate(tb.rows):
        cells = [c.text.strip()[:22] for c in row.cells]
        print(f"  r{ri}: {cells}")
