"""检查 docx 表格样式、标题run结构、正文run字体。"""
import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)

for i, t in enumerate(doc.tables):
    print(f"table{i}: style={t.style.name if t.style else None}")
    cell = t.rows[0].cells[0]
    for p in cell.paragraphs[:1]:
        for r in p.runs[:1]:
            f = r.font
            print(f"  cell run: font={f.name}, size={f.size}, bold={f.bold}")

print("--- headings runs ---")
for p in doc.paragraphs:
    if p.style and p.style.name.startswith("Heading"):
        runs_info = [(r.text[:30], r.font.name, r.font.size, r.font.bold) for r in p.runs]
        print(f"{p.text[:45]!r}: {runs_info}")
        if p.text.startswith("7. Practical"):
            break

print("--- body run sample (a Normal para) ---")
for p in doc.paragraphs:
    if p.text.startswith("The pattern confirms"):
        for r in p.runs:
            print(f"  run={r.text[:40]!r} font={r.font.name} size={r.font.size} bold={r.font.bold}")
        break
